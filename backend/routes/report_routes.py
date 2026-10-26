import os
import io
from bson import ObjectId
from flask import Blueprint, request, jsonify, send_file
from backend.db import get_db, serialize_doc
from backend.models.attendance_model import report_between
from backend.models.report_model import (
    add_report,
    get_report,
    list_reports,
    reports_for_person,
)
from datetime import date as date_type
from backend.utils.auth import admin_required, auth_required, current_auth

report_bp = Blueprint("reports", __name__)


def _to_oid(id_str):
    try:
        return ObjectId(id_str)
    except Exception:
        return id_str


def _person_from_auth():
    auth = current_auth()
    if not auth:
        return None, None
    role = auth.get("role")
    person_id = auth.get("sub")
    person_type = auth.get("person_type", "employee")
    if role == "admin":
        return "admin", person_id
    return person_type, person_id


@report_bp.route("/api/reports/export/<format_type>", methods=["GET"])
@admin_required
def export_report(format_type):
    report_type = request.args.get("type", "attendance")
    start_date = request.args.get("start_date")
    end_date = request.args.get("end_date")

    db = get_db()

    if report_type == "attendance":
        from backend.models.attendance_model import report_summary_between, auto_mark_absent_range
        if start_date and end_date:
            auto_mark_absent_range(start_date, end_date)
            data = report_summary_between(start_date, end_date)
        else:
            today = date_type.today()
            month_start = today.replace(day=1).isoformat()
            auto_mark_absent_range(month_start, today.isoformat())
            data = report_summary_between(month_start, today.isoformat())
        headers = ["Name", "Person Type", "ID", "Present", "Absent", "Late", "Total Hours"]
        rows = [
            [d.get("name", ""), d.get("person_type", ""), d.get("person_id", ""),
             d.get("present", 0), d.get("absent", 0), d.get("late", 0),
             d.get("total_hours", 0)]
            for d in data
        ]
    elif report_type == "employee":
        data = list(db.employees.find())
        headers = ["Employee ID", "Name", "Department", "Designation", "Phone", "Email", "Status"]
        rows = [
            [d.get("employee_id", ""), d.get("name", ""), d.get("department", ""),
             d.get("designation", ""), d.get("phone", ""), d.get("email", ""),
             d.get("status", "")]
            for d in data
        ]
    elif report_type == "intern":
        data = list(db.interns.find())
        headers = ["Intern ID", "Name", "College", "Department", "Mentor", "Duration"]
        rows = [
            [d.get("intern_id", ""), d.get("name", ""), d.get("college", ""),
             d.get("department", ""), d.get("mentor", ""), d.get("duration", "")]
            for d in data
        ]
    elif report_type == "task":
        query = {}
        if start_date and end_date:
            query["deadline"] = {"$gte": start_date, "$lte": end_date}
        data = list(db.tasks.find(query).sort("deadline", 1))
        headers = ["Title", "Assigned To", "Priority", "Deadline", "Status", "Remarks"]
        rows = [
            [d.get("title", ""), d.get("assigned_to", ""), d.get("priority", ""),
             d.get("deadline", ""), d.get("status", ""), d.get("remarks", "")]
            for d in data
        ]
    elif report_type == "finance":
        query = {}
        if start_date and end_date:
            query["date"] = {"$gte": start_date, "$lte": end_date}
        data = list(db.finance.find(query).sort("date", -1))
        headers = ["Date", "Type", "Category", "Description", "Amount", "Payment Method", "Status", "Notes"]
        rows = [
            [d.get("date", ""), d.get("type", ""), d.get("category", ""),
             d.get("description", ""), d.get("amount", ""), d.get("payment_method", ""),
             d.get("status", ""), d.get("notes", "")]
            for d in data
        ]
    else:
        return jsonify({"error": "Invalid report type"}), 400

    if format_type == "excel":
        return _export_excel(report_type, headers, rows)
    elif format_type == "pdf":
        return _export_pdf(report_type, headers, rows)
    else:
        return jsonify({"error": "Invalid format"}), 400


@report_bp.route("/api/reports", methods=["POST"])
@auth_required
def submit_report():
    """Submit (or update) today's work-progress report for the signed-in person."""
    person_type, person_id = _person_from_auth()
    if not person_id:
        return jsonify({"error": "Authentication required"}), 401

    data = request.get_json(silent=True) or {}
    date = data.get("date")
    if not date:
        from datetime import date as _date_type
        date = _date_type.today().isoformat()
    content = data.get("content", "")
    tasks_done = data.get("tasks_done")
    hours = data.get("hours")

    doc = add_report(person_type, person_id, date, content, tasks_done=tasks_done, hours=hours)
    if not doc:
        return jsonify({"error": "Could not save report"}), 500
    return jsonify({"message": "Report saved", "report": doc}), 200


@report_bp.route("/api/reports/mine", methods=["GET"])
@auth_required
def my_reports():
    """Reports submitted by the signed-in person."""
    person_type, person_id = _person_from_auth()
    if not person_id:
        return jsonify({"error": "Authentication required"}), 401
    docs = reports_for_person(person_id, person_type=person_type)
    return jsonify({"reports": docs}), 200


@report_bp.route("/api/reports", methods=["GET"])
@admin_required
def list_all_reports():
    """All daily reports (admin). Supports person_type, person_id, date,
    start_date and end_date query filters."""
    docs = list_reports(
        person_type=request.args.get("person_type"),
        person_id=request.args.get("person_id"),
        date=request.args.get("date"),
        start_date=request.args.get("start_date"),
        end_date=request.args.get("end_date"),
    )
    return jsonify({"reports": docs}), 200


@report_bp.route("/api/reports/<report_id>", methods=["GET"])
@auth_required
def get_one_report(report_id):
    person_type, person_id = _person_from_auth()
    if not person_id:
        return jsonify({"error": "Authentication required"}), 401
    doc = get_report(report_id)
    if not doc:
        return jsonify({"error": "Report not found"}), 404
    auth = current_auth()
    if auth.get("role") != "admin" and (
        doc.get("person_type") != person_type or doc.get("person_id") != person_id
    ):
        return jsonify({"error": "Forbidden"}), 403
    return jsonify({"report": doc}), 200


@report_bp.route("/api/reports/<report_id>", methods=["PUT"])
@auth_required
def update_report(report_id):
    person_type, person_id = _person_from_auth()
    if not person_id:
        return jsonify({"error": "Authentication required"}), 401
    doc = get_report(report_id)
    if not doc:
        return jsonify({"error": "Report not found"}), 404
    auth = current_auth()
    if auth.get("role") != "admin" and (
        doc.get("person_type") != person_type or doc.get("person_id") != person_id
    ):
        return jsonify({"error": "Forbidden"}), 403

    data = request.get_json(silent=True) or {}
    add_report(
        doc.get("person_type", person_type),
        doc.get("person_id", person_id),
        doc.get("date"),
        data.get("content", doc.get("content", "")),
        tasks_done=data.get("tasks_done", doc.get("tasks_done")),
        hours=data.get("hours", doc.get("hours")),
    )
    return jsonify({"message": "Report updated"}), 200


@report_bp.route("/api/reports/<report_id>", methods=["DELETE"])
@auth_required
def delete_report(report_id):
    person_type, person_id = _person_from_auth()
    if not person_id:
        return jsonify({"error": "Authentication required"}), 401
    db = get_db()
    doc = get_report(report_id)
    if not doc:
        return jsonify({"error": "Report not found"}), 404
    auth = current_auth()
    if auth.get("role") != "admin" and (
        doc.get("person_type") != person_type or doc.get("person_id") != person_id
    ):
        return jsonify({"error": "Forbidden"}), 403
    db.daily_reports.delete_one({"_id": _to_oid(report_id)})
    return jsonify({"message": "Report deleted"}), 200


def _export_excel(report_type, headers, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = report_type.title()

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            ws.cell(row=row_idx, column=col_idx, value=str(value) if value else "")

    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = max_length + 2

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"{report_type}_report.xlsx",
    )


def _export_pdf(report_type, headers, rows):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4))
    elements = []
    styles = getSampleStyleSheet()

    elements.append(Paragraph(f"{report_type.title()} Report", styles["Title"]))

    table_data = [headers] + [[str(v) if v else "" for v in row] for row in rows]

    table = Table(table_data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563eb")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 10),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
        ("BACKGROUND", (0, 1), (-1, -1), colors.white),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f0f4ff")]),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("BOX", (0, 0), (-1, -1), 1, colors.black),
    ]))

    elements.append(table)
    doc.build(elements)
    buffer.seek(0)

    return send_file(
        buffer,
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"{report_type}_report.pdf",
    )
